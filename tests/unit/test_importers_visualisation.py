"""Non-AGS importers (CSV, Excel, LAS, LandXML, ASCII grid, vector formats) and drawings
(sections, fence diagrams, HDD long sections, SVG / PDF)."""

import csv

import numpy as np
import pytest
from conftest import has_module

from agsgeomodel.importers import (
    read_ascii_grid,
    read_csv_project,
    read_dtm,
    read_excel,
    read_landxml,
    read_las,
    read_vector,
    tin_to_grid,
    to_alignment,
    to_faults,
    to_map_units,
)


def test_csv_project(tmp_path):
    loc = tmp_path / "locations.csv"
    ivs = tmp_path / "intervals.csv"
    with open(loc, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["Borehole", "Easting", "Northing", "Ground_Level", "Depth", "Type"])
        w.writerow(["BH1", "545000", "180000", "5.2", "20", "CP"])
        w.writerow(["TP1", "545050", "180010", "5.0", "3", "TP"])
    with open(ivs, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["borehole", "depth_from", "depth_to", "description", "code"])
        w.writerow(["BH1", "0", "1.5", "MADE GROUND", "MG"])
        w.writerow(["BH1", "1.5", "20", "Stiff grey CLAY", "LC"])
        w.writerow(["BH1", "5", "3", "bad interval", ""])  # base above top: skipped
    p = read_csv_project(loc, ivs)
    assert set(p.locations) == {"BH1", "TP1"} and p.locations["TP1"].kind.value == "TP"
    assert len(p.intervals) == 2 and p.locations["BH1"].ground_level == pytest.approx(5.2)


@pytest.mark.skipif(not has_module("openpyxl"), reason="openpyxl not installed")
def test_excel_flat_and_ags_layouts(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "locations"
    ws.append(["id", "x", "y", "gl", "depth"])
    ws.append(["BH1", 545000, 180000, 5.0, 10.0])
    ws2 = wb.create_sheet("intervals")
    ws2.append(["id", "top", "base", "description"])
    ws2.append(["BH1", 0, 10, "Stiff grey CLAY"])
    wb.save(tmp_path / "flat.xlsx")
    p = read_excel(tmp_path / "flat.xlsx")
    assert "BH1" in p.locations and len(p.intervals) == 1
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LOCA"
    ws.append(["HEADING", "LOCA_ID", "LOCA_NATE", "LOCA_NATN"])
    ws.append(["UNIT", "", "m", "m"])
    ws.append(["TYPE", "ID", "2DP", "2DP"])
    ws.append(["DATA", "BH1", "545000.00", "180000.00"])
    wb.save(tmp_path / "ags.xlsx")
    ags = read_excel(tmp_path / "ags.xlsx")
    assert ags["LOCA"].column("LOCA_ID") == ["BH1"] and ags["LOCA"].unit_of("LOCA_NATE") == "m"


def test_las_well_log(tmp_path):
    las = tmp_path / "w.las"
    las.write_text("""~VERSION INFORMATION
 VERS.   2.0 : CWLS LOG ASCII STANDARD - VERSION 2.0
 WRAP.   NO  : ONE LINE PER DEPTH STEP
~WELL INFORMATION
 STRT.M  1.0 : START DEPTH
 STOP.M  3.0 : STOP DEPTH
 STEP.M  1.0 : STEP
 NULL.   -999.25 : NULL VALUE
 WELL.   BH-LAS-1 : WELL
~CURVE INFORMATION
 DEPT.M        : DEPTH
 GR  .GAPI     : GAMMA RAY
~A
1.0  45.0
2.0  -999.25
3.0  80.0
""")
    w = read_las(las)
    assert w.well == "BH-LAS-1" and w.units["GR"] == "GAPI"
    np.testing.assert_allclose(w.curves["DEPT"], [1, 2, 3])
    assert np.isnan(w.curves["GR"][1])


def test_landxml_surface_and_alignment(tmp_path):
    x = tmp_path / "a.xml"
    x.write_text("""<?xml version="1.0"?>
<LandXML xmlns="http://www.landxml.org/schema/LandXML-1.2" version="1.2">
 <Surfaces><Surface name="EG"><Definition surfType="TIN"><Pnts>
  <P id="1">180000 545000 5.0</P><P id="2">180000 545100 6.0</P><P id="3">180100 545000 7.0</P>
  <P id="4">180100 545100 8.0</P></Pnts>
  <Faces><F>1 2 3</F><F>2 4 3</F></Faces></Definition></Surface></Surfaces>
 <Alignments><Alignment name="HDD" length="100"><CoordGeom>
  <Line><Start>180000 545000</Start><End>180000 545050</End></Line>
  <Line><Start>180000 545050</Start><End>180000 545100</End></Line></CoordGeom>
  <Profile><ProfAlign><PVI>0 5</PVI><PVI>50 -5</PVI><PVI>100 6</PVI></ProfAlign></Profile>
 </Alignment></Alignments>
</LandXML>""")
    d = read_landxml(x)
    V, F = d["surfaces"]["EG"]
    assert V.shape == (4, 3) and F.shape == (2, 3) and V[0, 0] == 545000.0  # easting first
    g, z = tin_to_grid(V, F, 10.0)
    assert np.nanmin(z) >= 5.0 - 1e-9 and np.nanmax(z) <= 8.0 + 1e-9
    al = d["alignments"]["HDD"]
    np.testing.assert_allclose(al["line"], [[545000, 180000], [545050, 180000], [545100, 180000]])
    assert al["profile"].shape == (3, 2)


def test_ascii_grid_and_dtm(example_inputs):
    out, _ = example_inputs
    g, v = read_ascii_grid(out / "dtm_5m.asc")
    g2, v2 = read_dtm(out / "dtm_5m.asc")
    assert g == g2 and v.shape == (g.ny, g.nx) and np.isfinite(v).mean() > 0.99


def test_vector_conversions(example_inputs):
    out, _ = example_inputs
    faults = to_faults(read_vector(out / "faults.geojson"))
    assert faults and faults[0].downthrown_side == "right" and faults[0].throw_m is None
    line, props = to_alignment(read_vector(out / "hdd_alignment.geojson"))
    assert line.shape[1] == 2 and props["name"]
    polys = [
        {
            "geometry": {"type": "Polygon", "coordinates": [[[0, 0], [10, 0], [10, 10], [0, 0]]]},
            "properties": {"LEX": "ALV"},
        }
    ]
    mu = to_map_units(polys, "LEX", "superficial", default_scale=50000)
    assert mu[0].code == "ALV" and mu[0].nominal_scale == 50000


# ---------------------------------------------------------------------------- drawings
def test_section_svg_pdf(model, interpreted, tmp_path):
    from agsgeomodel.examples.synthetic import E0, N0
    from agsgeomodel.visualisation import compute_section, render_section

    p, it = interpreted
    sec = compute_section(
        model, np.array([[E0 + 100, N0 + 300], [E0 + 1150, N0 + 300]]), name="T", project=p, interpretation=it
    )
    assert sec.logs and sec.faults  # holes projected and the fault crossed
    for u in sec.units:
        ok = np.isfinite(sec.tops[u]) & np.isfinite(sec.bases[u])
        assert np.all(sec.tops[u][ok] >= sec.bases[u][ok] - 1e-9)
    d = render_section(sec, it.column.kb, vertical_exaggeration=5.0)
    d.save_svg(tmp_path / "s.svg")
    d.save_pdf(tmp_path / "s.pdf")
    svg = (tmp_path / "s.svg").read_text()
    assert svg.startswith(("<?xml", "<svg"))
    assert "London Clay Formation" in svg and "Speculative" in svg  # legend: units and confidence classes
    pdf = (tmp_path / "s.pdf").read_bytes()
    assert pdf.startswith(b"%PDF-") and pdf.rstrip().endswith(b"%%EOF")
    if has_module("pypdf"):
        import pypdf

        assert len(pypdf.PdfReader(tmp_path / "s.pdf").pages) == 1


def test_fence_and_hdd_long_section(example_project, tmp_path):
    from agsgeomodel.hdd import render_hdd_long_section
    from agsgeomodel.visualisation import render_fence_2d

    _, res = example_project
    assert res.fence
    render_fence_2d(res.fence, res.interpretation.column.kb).save_svg(tmp_path / "f.svg")
    assert (tmp_path / "f.svg").stat().st_size > 5000
    d = render_hdd_long_section(res.hdd, res.model, res.project, res.interpretation, res.hydro)
    d.save_pdf(tmp_path / "h.pdf")
    svg = d.to_svg()
    for word in ("running sand", "groundwater", "mixed face"):
        assert word in svg.lower()
