"""GIS and tabular exports: GeoJSON, ESRI Shapefile (pure Python), CSV, Excel, PNG rasters.

Features are GeoJSON-like dicts ``{"geometry": {...}, "properties": {...}}`` so any table from the
ProjectStore (or computed in memory) can be exported in any format.
"""

from __future__ import annotations

import csv
import datetime as _dt
import json
import struct
import zlib
from pathlib import Path

import numpy as np

from ..core.optional import require
from ..database.gpkg import KNOWN_WKT


def _clean(v):
    if isinstance(v, (np.floating, np.integer)):
        v = v.item()
    if isinstance(v, float) and v != v:
        return None
    if isinstance(v, (dict, list, tuple)):
        return json.dumps(v, default=str)
    return v


# ============================================================================ GeoJSON
def write_geojson(path, features: list[dict], crs: str | None = None, name: str | None = None) -> Path:
    """GeoJSON FeatureCollection. Coordinates stay in the project CRS; a legacy ``crs`` member is
    written so QGIS/GDAL read it correctly (RFC 7946 assumes WGS84 — transform first if required)."""
    path = Path(path)
    fc = {
        "type": "FeatureCollection",
        "name": name or path.stem,
        "features": [
            {
                "type": "Feature",
                "geometry": f.get("geometry"),
                "properties": {k: _clean(v) for k, v in f.get("properties", {}).items()},
            }
            for f in features
        ],
    }
    if crs:
        code = str(crs).upper().replace("EPSG:", "")
        fc["crs"] = {"type": "name", "properties": {"name": f"urn:ogc:def:crs:EPSG::{code}"}}
    path.write_text(json.dumps(fc, default=str))
    return path


# ============================================================================ Shapefile
_SHP_TYPES = {
    ("Point", False): 1,
    ("LineString", False): 3,
    ("Polygon", False): 5,
    ("MultiPoint", False): 8,
    ("Point", True): 11,
    ("LineString", True): 13,
    ("Polygon", True): 15,
    ("MultiPoint", True): 18,
    ("MultiLineString", False): 3,
    ("MultiLineString", True): 13,
    ("MultiPolygon", False): 5,
    ("MultiPolygon", True): 15,
}


def _parts(geom: dict) -> list[list]:
    t, c = geom["type"], geom["coordinates"]
    if t == "LineString":
        return [c]
    if t == "MultiLineString":
        return list(c)
    if t == "Polygon":
        return list(c)
    if t == "MultiPolygon":
        return [r for poly in c for r in poly]
    if t == "MultiPoint":
        return [c]
    raise ValueError(t)


def _ring_cw(ring) -> list:
    a = np.asarray(ring, float)
    area = 0.5 * float(np.dot(a[:, 0], np.roll(a[:, 1], -1)) - np.dot(a[:, 1], np.roll(a[:, 0], -1)))
    return list(ring) if area <= 0 else list(ring)[::-1]


def write_shapefile(path, features: list[dict], crs: str | None = None, encoding: str = "UTF-8") -> Path:
    """Write .shp/.shx/.dbf/.prj/.cpg. Field names are truncated to 10 characters (made unique);
    polygon exteriors are written clockwise and holes counter-clockwise as the format requires."""
    path = Path(path).with_suffix(".shp")
    geoms = [f.get("geometry") for f in features]
    gtypes = {g["type"] for g in geoms if g}
    if not gtypes:
        raise ValueError("no geometries")
    base_t = next(iter(gtypes)).replace("Multi", "") if len({t.replace("Multi", "") for t in gtypes}) == 1 else None
    if base_t is None:
        raise ValueError(f"mixed geometry types {gtypes}")
    if "MultiPoint" in gtypes:
        raise ValueError("MultiPoint shapefiles are not supported; explode to points first")
    has_z = any(len(_first(g)) > 2 for g in geoms if g)
    gt = next(iter(gtypes))
    stype = _SHP_TYPES[(gt if gt in ("Point", "MultiPoint") else base_t, has_z)]
    recs: list[bytes] = []
    bbox = [np.inf, np.inf, -np.inf, -np.inf]
    zr = [np.inf, -np.inf]
    for k, g in enumerate(geoms, start=1):
        if g is None:
            content = struct.pack("<i", 0)
        elif stype in (1, 11):
            x, y = g["coordinates"][:2]
            z = g["coordinates"][2] if len(g["coordinates"]) > 2 else 0.0
            content = struct.pack("<i2d", stype, x, y)
            if stype == 11:
                content += struct.pack("<2d", z, 0.0)
            bbox = [min(bbox[0], x), min(bbox[1], y), max(bbox[2], x), max(bbox[3], y)]
            zr = [min(zr[0], z), max(zr[1], z)]
        else:
            parts = _parts(g)
            if stype in (5, 15):
                fixed = []
                polys = [g["coordinates"]] if g["type"] == "Polygon" else g["coordinates"]
                for poly in polys:
                    fixed.append(_ring_cw(poly[0]))
                    fixed.extend(_ring_cw(h)[::-1] for h in poly[1:])
                parts = fixed
            pts = [p for part in parts for p in part]
            arr = np.array([[p[0], p[1]] for p in pts], float)
            zs = np.array([p[2] if len(p) > 2 else 0.0 for p in pts], float)
            bb = [arr[:, 0].min(), arr[:, 1].min(), arr[:, 0].max(), arr[:, 1].max()]
            bbox = [min(bbox[0], bb[0]), min(bbox[1], bb[1]), max(bbox[2], bb[2]), max(bbox[3], bb[3])]
            starts, n = [], 0
            for part in parts:
                starts.append(n)
                n += len(part)
            content = struct.pack("<i4d2i", stype, *bb, len(parts), len(pts))
            content += struct.pack(f"<{len(parts)}i", *starts)
            content += arr.astype("<f8").tobytes()
            if stype in (13, 15, 18):
                content += struct.pack("<2d", zs.min(), zs.max()) + zs.astype("<f8").tobytes()
                content += struct.pack("<2d", 0.0, 0.0) + np.zeros(len(pts)).astype("<f8").tobytes()
                zr = [min(zr[0], zs.min()), max(zr[1], zs.max())]
        recs.append(struct.pack(">2i", k, len(content) // 2) + content)
    if not np.isfinite(bbox[0]):
        bbox = [0, 0, 0, 0]
    if not np.isfinite(zr[0]):
        zr = [0, 0]

    def header(file_len_words: int) -> bytes:
        return (
            struct.pack(">7i", 9994, 0, 0, 0, 0, 0, file_len_words)
            + struct.pack("<2i", 1000, stype)
            + struct.pack("<4d", *bbox)
            + struct.pack("<4d", zr[0], zr[1], 0.0, 0.0)
        )

    shp_body = b"".join(recs)
    path.write_bytes(header((100 + len(shp_body)) // 2) + shp_body)
    shx = bytearray()
    off = 100
    for r in recs:
        shx += struct.pack(">2i", off // 2, (len(r) - 8) // 2)
        off += len(r)
    path.with_suffix(".shx").write_bytes(header((100 + len(shx)) // 2) + bytes(shx))
    _write_dbf(path.with_suffix(".dbf"), [f.get("properties", {}) for f in features], encoding)
    path.with_suffix(".cpg").write_text(encoding)
    if crs:
        code = int(str(crs).upper().replace("EPSG:", ""))
        wkt = KNOWN_WKT.get(code)
        if wkt is None:
            try:
                import pyproj  # type: ignore

                wkt = pyproj.CRS.from_epsg(code).to_wkt("WKT1_ESRI")
            except Exception:
                wkt = None
        if wkt:
            path.with_suffix(".prj").write_text(wkt)
    return path


def _first(g) -> list:
    c = g["coordinates"]
    while isinstance(c, (list, tuple)) and c and isinstance(c[0], (list, tuple)):
        c = c[0]
    return list(c)


def _write_dbf(path: Path, rows: list[dict], encoding: str) -> None:
    keys: list[str] = []
    for r in rows:
        for k in r:
            if k not in keys and k not in ("geometry",):
                keys.append(k)
    names, used = [], set()
    for k in keys:
        n = "".join(ch if ch.isalnum() or ch == "_" else "_" for ch in k)[:10] or "F"
        base, i = n, 1
        while n.upper() in used:
            suf = str(i)
            n = base[: 10 - len(suf)] + suf
            i += 1
        used.add(n.upper())
        names.append(n)
    specs = []
    for k in keys:
        vals = [_clean(r.get(k)) for r in rows]
        nonnull = [v for v in vals if v is not None]
        if nonnull and all(isinstance(v, bool) for v in nonnull):
            specs.append(("L", 1, 0))
        elif nonnull and all(isinstance(v, int) and not isinstance(v, bool) for v in nonnull):
            specs.append(("N", 18, 0))
        elif nonnull and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in nonnull):
            specs.append(("N", 19, 6))
        else:
            width = min(254, max([len(str(v).encode(encoding, "replace")) for v in nonnull] + [1]))
            specs.append(("C", width, 0))
    rec_len = 1 + sum(s[1] for s in specs)
    today = _dt.date.today()
    head = struct.pack(
        "<BBBBIHH20x", 3, today.year - 1900, today.month, today.day, len(rows), 32 + 32 * len(keys) + 1, rec_len
    )
    fields = b""
    for n, (t, w, d) in zip(names, specs):
        fields += struct.pack("<11sc4xBB14x", n.encode("ascii", "replace"), t.encode(), w, d)
    out = bytearray(head + fields + b"\r")
    for r in rows:
        out += b" "
        for k, (t, w, d) in zip(keys, specs):
            v = _clean(r.get(k))
            if v is None:
                s = b" " * w
            elif t == "L":
                s = b"T" if v else b"F"
            elif t == "N":
                txt = f"{v:.{d}f}" if d else str(int(v))
                s = txt.encode()[:w].rjust(w)
            else:
                s = str(v).encode(encoding, "replace")[:w].ljust(w)
            out += s
    out += b"\x1a"
    path.write_bytes(bytes(out))


# ============================================================================ tables
def write_csv(path, rows: list[dict]) -> Path:
    path = Path(path)
    keys: list[str] = []
    for r in rows:
        for k in r:
            if k not in keys and k != "geometry":
                keys.append(k)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=keys, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: _clean(r.get(k)) for k in keys})
    return path


def write_excel(path, sheets: dict[str, list[dict]], notes: dict[str, str] | None = None) -> Path:
    """Multi-sheet workbook (requires openpyxl). A README sheet carries provenance notes."""
    openpyxl = require("openpyxl", "Excel export")
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "README"
    ws.append(["AGSGeoModel export", _dt.datetime.now().isoformat(timespec="seconds")])
    for k, v in (notes or {}).items():
        ws.append([k, v])
    for name, rows in sheets.items():
        sh = wb.create_sheet(name[:31])
        keys: list[str] = []
        for r in rows:
            for k in r:
                if k not in keys and k != "geometry":
                    keys.append(k)
        sh.append(keys)
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="37474F")
        for r in rows:
            sh.append([_clean(r.get(k)) for k in keys])
        sh.freeze_panes = "A2"
        for col in sh.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
            sh.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
    path = Path(path)
    wb.save(path)
    return path


# ============================================================================ PNG
def write_png(path, rgba: np.ndarray) -> Path:
    """Write an (H, W, 3|4) uint8 array as PNG (zlib only)."""
    a = np.asarray(rgba, np.uint8)
    if a.ndim == 2:
        a = np.stack([a] * 3, axis=-1)
    h, w, c = a.shape
    ctype = {3: 2, 4: 6}[c]
    raw = b"".join(b"\x00" + a[y].tobytes() for y in range(h))

    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    png = (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, ctype, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw, 6))
        + chunk(b"IEND", b"")
    )
    path = Path(path)
    path.write_bytes(png)
    return path


def grid_to_rgba(
    values: np.ndarray, ramp: list[str] | None = None, vmin=None, vmax=None, categorical: dict[int, str] | None = None
) -> np.ndarray:
    v = np.asarray(values, float)[::-1]
    out = np.zeros(v.shape + (4,), np.uint8)
    ok = np.isfinite(v)

    def rgb(h):
        h = h.lstrip("#")
        return [int(h[i : i + 2], 16) for i in (0, 2, 4)]

    if categorical:
        for code, colour in categorical.items():
            m = ok & (np.rint(v) == code)
            out[m, :3] = rgb(colour)
            out[m, 3] = 255
        return out
    ramp = ramp or ["#2C7BB6", "#ABD9E9", "#FFFFBF", "#FDAE61", "#D7191C"]
    lo = np.nanmin(v) if vmin is None else vmin
    hi = np.nanmax(v) if vmax is None else vmax
    t = np.clip((v - lo) / max(hi - lo, 1e-12), 0, 1)
    stops = np.linspace(0, 1, len(ramp))
    cols = np.array([rgb(c) for c in ramp], float)
    for ch in range(3):
        out[..., ch] = np.interp(t, stops, cols[:, ch]).astype(np.uint8)
    out[..., 3] = np.where(ok, 255, 0)
    return out


def drawing_to_png(drawing, path, dpi: int = 150) -> Path:
    """Rasterise an SVG drawing: cairosvg if installed, else Qt (inside QGIS)."""
    path = Path(path)
    svg = drawing.to_svg().encode()
    try:
        import cairosvg  # type: ignore

        cairosvg.svg2png(bytestring=svg, write_to=str(path), dpi=dpi)
        return path
    except ImportError:
        pass
    try:
        from qgis.PyQt.QtCore import QByteArray, QSize  # type: ignore
        from qgis.PyQt.QtGui import QColor, QImage, QPainter  # type: ignore
        from qgis.PyQt.QtSvg import QSvgRenderer  # type: ignore

        w = int(drawing.width / 25.4 * dpi)
        h = int(drawing.height / 25.4 * dpi)
        img = QImage(QSize(w, h), QImage.Format.Format_ARGB32)
        img.fill(QColor(255, 255, 255))
        r = QSvgRenderer(QByteArray(svg))
        p = QPainter(img)
        r.render(p)
        p.end()
        img.save(str(path))
        return path
    except ImportError as exc:  # pragma: no cover - environment dependent
        from ..core.errors import DependencyMissing

        raise DependencyMissing("cairosvg", "PNG export of drawings outside QGIS") from exc
